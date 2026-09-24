import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation_alt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette 2: Warm Professional & Modern (Dark Indigo & Amber)
    c_dark = RGBColor(30, 27, 75)      # Deep Indigo 950
    c_accent = RGBColor(245, 158, 11)  # Amber 500
    c_white = RGBColor(255, 255, 255)
    c_gray = RGBColor(71, 85, 105)     # Slate 600
    c_card_bg = RGBColor(255, 255, 255)
    c_card_border = RGBColor(226, 232, 240)
    c_light_bg = RGBColor(248, 250, 252)

    def add_header(slide, title_text, subtitle_text=""):
        # Top banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
        banner.fill.solid()
        banner.fill.fore_color.rgb = c_dark
        banner.line.fill.background()

        # Accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.2), prs.slide_width, Inches(0.08))
        line.fill.solid()
        line.fill.fore_color.rgb = c_accent
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = c_white

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Arial"
            p2.font.size = Pt(13)
            p2.font.color.rgb = c_accent
            p2.space_before = Pt(2)

    # ==================== SLIDE 1: PORTADA ALTERNATIVA ====================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_dark
    bg1.line.fill.background()

    # Decorative geometric element
    dec = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    dec.fill.solid()
    dec.fill.fore_color.rgb = RGBColor(49, 46, 129) # Indigo 900
    dec.line.color.rgb = c_accent
    dec.line.width = Pt(1.5)

    tb1 = slide1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Calidad del Software & Ecosistema de IA"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_white

    p_sub = tf1.add_paragraph()
    p_sub.text = "Enfoque Práctico para el Desarrollo de Software (ADSO - SENA)"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = c_accent
    p_sub.space_before = Pt(12)

    p_info = tf1.add_paragraph()
    p_info.text = "• Expositores: Yedinson Ortiz Pino, William Oswaldo Luna, Alejandro Euscategui\n• Ficha: 3223899 | Centro de Comercio y Servicios\n• Fecha: 19 de agosto de 2026"
    p_info.font.name = "Arial"
    p_info.font.size = Pt(14)
    p_info.font.color.rgb = RGBColor(226, 232, 240)
    p_info.space_before = Pt(28)

    # ==================== SLIDE 2: CALIDAD VS FUNCIONALIDAD ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Calidad vs. Funcionalidad Básica", "¿Por qué un software funcional no siempre es un software de calidad?")

    # 3 Column layout
    cols = [
        ("1. Funcionalidad", "El programa cumple con las tareas principales solicitadas, ejecuta código y da resultados aparentes.", RGBColor(238, 242, 255)),
        ("2. Riesgos Ocultos", "Vulnerabilidades de seguridad, fugas de memoria (Memory Leaks), deuda técnica y código espagueti.", RGBColor(254, 242, 242)),
        ("3. Calidad Real", "Mantenible en el tiempo, seguro ante ataques, eficiente en recursos y altamente usable para el usuario final.", RGBColor(240, 253, 244))
    ]

    for i, (title, desc, color) in enumerate(cols):
        left = Inches(0.8 + i * 4.0)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = color
        card.line.color.rgb = c_card_border
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.4)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = c_dark

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(14)
        p2.font.color.rgb = c_gray
        p2.space_before = Pt(16)

    # ==================== SLIDE 3: CLASIFICACIÓN DE HERRAMIENTAS DE IA ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Ecosistema de Asistentes de IA", "Clasificación de herramientas según su integración en el desarrollo")

    categories = [
        ("Integradas en el IDE", "GitHub Copilot & Cursor", "Asistencia en tiempo real, autocompletado y refactorización directa en el editor de código."),
        ("Asistentes Conversacionales", "ChatGPT, Claude & Gemini", "Ideales para explicar conceptos, resolver dudas de arquitectura y redactar scripts generales."),
        ("Especializadas y de Nicho", "Astral, Grok, OpenCode & Stitch", "Optimización de sintaxis específica, análisis de rendimiento y prototipado rápido de interfaces web.")
    ]

    for i, (cat, tools, desc) in enumerate(categories):
        top = Inches(1.8 + i * 1.7)
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = c_light_bg
        box.line.color.rgb = c_card_border
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.4)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = cat + " — "
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_dark

        run = p.add_run()
        run.text = tools
        run.font.bold = True
        run.font.color.rgb = c_accent

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_gray
        p2.space_before = Pt(4)

    # ==================== SLIDE 4: CONCLUSIÓN ADSO ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Conclusión: Elección Estratégica para ADSO", "Por qué GitHub Copilot destaca para estudiantes de programación")

    box_main = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    box_main.fill.solid()
    box_main.fill.fore_color.rgb = c_light_bg
    box_main.line.color.rgb = c_card_border
    tf_bm = box_main.text_frame
    tf_bm.word_wrap = True
    tf_bm.margin_left = tf_bm.margin_right = tf_bm.margin_top = tf_bm.margin_bottom = Inches(0.5)

    p = tf_bm.paragraphs[0]
    p.text = "Recomendación Principal: GitHub Copilot"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = c_dark

    points = [
        "• Flujo de trabajo optimizado: Funciona nativamente como extensión en VS Code, evitando cambios de contexto constantes.",
        "• Accesibilidad económica: Totalmente gratuito para estudiantes mediante el paquete oficial de GitHub.",
        "• Cobertura tecnológica: Excelente rendimiento en lenguajes y frameworks estudiados en el SENA (PHP, Laravel, JS, HTML, CSS, SQL).",
        "• Impulso a la calidad y pruebas: Agiliza la generación de comentarios, documentación y pruebas unitarias.",
        "• Aprendizaje activo: Obliga al aprendiz a analizar, probar y validar la lógica sugerida antes de integrarla al proyecto."
    ]

    for pt in points:
        pp = tf_bm.add_paragraph()
        pp.text = pt
        pp.font.name = "Arial"
        pp.font.size = Pt(14)
        pp.font.color.rgb = c_gray
        pp.space_before = Pt(10)

    output_path = "C:/Users/Personal/Documents/Default Project/Calidad_Software_IA_Alternativa.pptx"
    prs.save(output_path)
    print(f"Presentación alternativa guardada en: {output_path}")

if __name__ == "__main__":
    create_presentation_alt()
