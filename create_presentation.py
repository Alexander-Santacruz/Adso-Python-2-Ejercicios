import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_dark = RGBColor(15, 23, 42)      # #0F172A Slate 900
    c_green = RGBColor(16, 185, 129)   # #10B981 Emerald 500
    c_white = RGBColor(255, 255, 255)
    c_gray = RGBColor(100, 116, 139)   # #64748B Slate 500
    c_card_bg = RGBColor(248, 250, 252)# #F8FAFC Slate 50
    c_card_border = RGBColor(226, 232, 240)

    def add_header(slide, title_text, subtitle_text=""):
        # Header banner or clean top text
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = c_dark

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Arial"
            p2.font.size = Pt(14)
            p2.font.color.rgb = c_green
            p2.space_before = Pt(4)

    # ==================== SLIDE 1: PORTADA ====================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_dark
    bg1.line.fill.background()

    # Accent decorative bar
    bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.15), Inches(3.8))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = c_green
    bar1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.6), Inches(1.7), Inches(10.5), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Calidad del Software e Inteligencia Artificial"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_white

    p_sub = tf1.add_paragraph()
    p_sub.text = "SENA — Análisis y Desarrollo de Software (ADSO)"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = c_green
    p_sub.space_before = Pt(12)

    p_info = tf1.add_paragraph()
    p_info.text = "Integrantes: Yedinson Ortiz Pino, William Oswaldo Luna, Alejandro Euscategui\nFicha: 3223899 | Centro de Comercio y Servicios\nFecha: 19 de agosto de 2026"
    p_info.font.name = "Arial"
    p_info.font.size = Pt(14)
    p_info.font.color.rgb = RGBColor(203, 213, 225)
    p_info.space_before = Pt(24)

    # ==================== SLIDE 2: QUÉ ES SOFTWARE Y CALIDAD ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Conceptos Fundamentales", "Definición y pilares de la calidad del software")

    # Card 1: Software
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5))
    c1.fill.solid()
    c1.fill.fore_color.rgb = c_card_bg
    c1.line.color.rgb = c_card_border
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_right = tf_c1.margin_top = tf_c1.margin_bottom = Inches(0.4)
    
    p = tf_c1.paragraphs[0]
    p.text = "¿Qué es Software?"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = c_dark

    p2 = tf_c1.add_paragraph()
    p2.text = "Conjunto de programas, instrucciones, reglas informáticas, datos y documentación asociada que permiten a un sistema informático ejecutar tareas específicas."
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = c_gray
    p2.space_before = Pt(12)

    p3 = tf_c1.add_paragraph()
    p3.text = "• Código fuente y ejecutables\n• Configuraciones y bases de datos\n• Manuales de usuario"
    p3.font.name = "Arial"
    p3.font.size = Pt(13)
    p3.font.color.rgb = c_dark
    p3.space_before = Pt(12)

    # Card 2: Calidad de Software
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.5))
    c2.fill.solid()
    c2.fill.fore_color.rgb = c_card_bg
    c2.line.color.rgb = c_card_border
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_right = tf_c2.margin_top = tf_c2.margin_bottom = Inches(0.4)

    p = tf_c2.paragraphs[0]
    p.text = "¿Qué es Buena Calidad?"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = c_dark

    p2 = tf_c2.add_paragraph()
    p2.text = "No solo cumple las funciones para las que fue diseñado, sino que destaca en estándares técnicos y experiencia de usuario:"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = c_gray
    p2.space_before = Pt(12)

    p3 = tf_c2.add_paragraph()
    p3.text = "• Seguridad y confiabilidad\n• Eficiencia y optimización de recursos\n• Usabilidad (fácil de usar)\n• Mantenibilidad en el tiempo y libre de fallos"
    p3.font.name = "Arial"
    p3.font.size = Pt(13)
    p3.font.color.rgb = c_dark
    p3.space_before = Pt(12)

    # ==================== SLIDE 3: EL MITO DEL "FUNCIONA CORRECTAMENTE" ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "¿Funcionar correctamente es sinónimo de calidad?", "La diferencia entre que 'compile' y que sea un software robusto")

    # Left: Myth
    box_m = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5))
    box_m.fill.solid()
    box_m.fill.fore_color.rgb = RGBColor(254, 242, 242) # soft red
    box_m.line.color.rgb = RGBColor(254, 202, 202)
    tf_m = box_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_right = tf_m.margin_top = tf_m.margin_bottom = Inches(0.4)

    p = tf_m.paragraphs[0]
    p.text = "❌ El Mito"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(185, 28, 28)

    p2 = tf_m.add_paragraph()
    p2.text = "\"Si un programa ejecuta sus tareas principales (funcionalidad) de manera correcta, ya se considera software de calidad.\""
    p2.font.name = "Arial"
    p2.font.size = Pt(15)
    p2.font.color.rgb = c_dark
    p2.space_before = Pt(16)

    # Right: Reality (Errors hidden)
    box_r = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.5))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = c_card_bg
    box_r.line.color.rgb = c_card_border
    tf_r = box_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = Inches(0.4)

    p = tf_r.paragraphs[0]
    p.text = "⚠️ Errores ocultos a pesar de funcionar:"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = c_dark

    p2 = tf_r.add_paragraph()
    p2.text = "• Vulnerabilidades de seguridad (Inyección SQL, datos expuestos)\n• Fugas de memoria (Memory Leaks) que degradan el sistema\n• Errores lógicos en casos borde (Edge Cases)\n• Deuda técnica y código espagueti\n• Incompatibilidad de entorno o baja escalabilidad"
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_gray
    p2.space_before = Pt(12)

    # ==================== SLIDE 4: HERRAMIENTAS DE IA EN EL DESARROLLO ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Herramientas de Inteligencia Artificial", "Asistentes modernos para optimizar el ciclo de desarrollo")

    ai_tools = [
        ("GitHub Copilot", "IDE (VS Code)", "Autocompletado en tiempo real mientras escribes código.", "Gratis (Estudiantes)"),
        ("ChatGPT / GPT-4", "Web / API", "Explicación de conceptos, solución de errores y lógica paso a paso.", "Gratis / Pago"),
        ("Claude (Anthropic)", "Web / Cursor", "Análisis profundo de contexto, refactorización y revisión de seguridad.", "Gratis / Pago"),
        ("Gemini (Google)", "Web / IDE", "Integración con Google, soporte multimodal y análisis de arquitectura.", "Gratis / Pago"),
        ("Cursor", "IDE Completo", "Editor impulsado por IA para generar y modificar proyectos completos.", "Gratis / Pago"),
        ("OpenCode / Otros", "CLI / Local", "Modelos abiertos y herramientas de automatización de sintaxis.", "Open Source")
    ]

    for i, (name, platform, desc, pricing) in enumerate(ai_tools):
        col = i % 3
        row = i // 3
        left = Inches(0.8 + col * 3.95)
        top = Inches(2.0 + row * 2.5)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.75), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_card_border
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = name
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_dark

        p_plat = tf.add_paragraph()
        p_plat.text = f"Plataforma: {platform} | {pricing}"
        p_plat.font.name = "Arial"
        p_plat.font.size = Pt(11)
        p_plat.font.color.rgb = c_green
        p_plat.space_before = Pt(2)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = c_gray
        p_desc.space_before = Pt(6)

    # ==================== SLIDE 5: ANÁLISIS Y CONCLUSIÓN (ADSO) ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Herramienta Recomendada para ADSO", "¿Por qué GitHub Copilot es la opción ideal para estudiantes?")

    # Left box: Selected tool highlight
    box_left = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(4.0), Inches(4.5))
    box_left.fill.solid()
    box_left.fill.fore_color.rgb = c_dark
    box_left.line.fill.background()
    tf_bl = box_left.text_frame
    tf_bl.word_wrap = True
    tf_bl.margin_left = tf_bl.margin_right = tf_bl.margin_top = tf_bl.margin_bottom = Inches(0.4)

    p = tf_bl.paragraphs[0]
    p.text = "🏆 GitHub Copilot"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = c_white

    p2 = tf_bl.add_paragraph()
    p2.text = "La elección perfecta para los aprendices del programa Análisis y Desarrollo de Software (SENA)."
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_before = Pt(16)

    # Right box: 5 Arguments
    box_right = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(2.0), Inches(7.4), Inches(4.5))
    box_right.fill.solid()
    box_right.fill.fore_color.rgb = c_card_bg
    box_right.line.color.rgb = c_card_border
    tf_br = box_right.text_frame
    tf_br.word_wrap = True
    tf_br.margin_left = tf_br.margin_right = tf_br.margin_top = tf_br.margin_bottom = Inches(0.4)

    p = tf_br.paragraphs[0]
    p.text = "5 Argumentos Clave:"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_dark

    args = [
        "1. Integración directa en VS Code (sin cambiar de ventana).",
        "2. Gratis para estudiantes mediante el GitHub Student Developer Pack.",
        "3. Soporte total para tecnologías SENA: PHP, Laravel, JS, HTML, CSS, SQL.",
        "4. Facilita la documentación y creación de pruebas unitarias.",
        "5. Fomenta el aprendizaje crítico al obligar a leer y verificar el código sugerido."
    ]

    for arg in args:
        pa = tf_br.add_paragraph()
        pa.text = arg
        pa.font.name = "Arial"
        pa.font.size = Pt(13)
        pa.font.color.rgb = c_gray
        pa.space_before = Pt(8)

    output_path = "C:/Users/Personal/Documents/Default Project/Calidad_Software_IA.pptx"
    prs.save(output_path)
    print(f"Presentación guardada exitosamente en: {output_path}")

if __name__ == "__main__":
    create_presentation()
